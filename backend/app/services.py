import os
import json
import httpx
import pandas as pd
from typing import List, Optional, Dict, Any, Union
from sqlalchemy.orm import Session
from .models import User, Group, Prompt, Record, Knowledge, Task, History, Matrix
from .schemas import KnowledgeMatrixItem, MatrixCreate, MatrixUpdate, MatrixExportData
from datetime import datetime
import asyncio

class AIService:
    def __init__(self):
        self.openai_api_key = os.getenv("OPENAI_API_KEY")
        self.anthropic_api_key = os.getenv("ANTHROPIC_API_KEY")
        self.gemini_api_key = os.getenv("GEMINI_API_KEY")
        self.lm_studio_base_url = os.getenv("LM_STUDIO_BASE_URL", "http://localhost:1234/v1")
        self.default_provider = os.getenv("DEFAULT_AI_PROVIDER", "openai")
        self.default_model = os.getenv("DEFAULT_AI_MODEL", "gpt-4o-mini")

    async def generate_knowledge(self, content: str, prompt: str, provider: Optional[str] = None, model: Optional[str] = None) -> Dict[str, str]:
        """Generate knowledge using AI provider"""
        provider = provider or self.default_provider
        model = model or self.default_model
        
        if self._is_test_mode():
            return self._generate_test_response(content, prompt)
        
        try:
            if provider == "openai":
                return await self._call_openai(content, prompt, model)
            elif provider == "anthropic":
                return await self._call_anthropic(content, prompt, model)
            elif provider == "gemini":
                return await self._call_gemini(content, prompt, model)
            elif provider == "lmstudio":
                return await self._call_lmstudio(content, prompt, model)
            else:
                raise ValueError(f"Unsupported AI provider: {provider}")
        except Exception as e:
            print(f"AI generation failed: {str(e)}, using test response")
            return self._generate_test_response(content, prompt)

    async def _call_openai(self, content: str, prompt: str, model: str) -> Dict[str, str]:
        headers = {
            "Authorization": f"Bearer {self.openai_api_key}",
            "Content-Type": "application/json"
        }
        
        # task2knowledgeと同じようにプロンプトとコンテンツを結合
        request_content = f"{prompt}\n\n{content}"
        
        data = {
            "model": model,
            "messages": [  
                {"role": "user", "content": request_content}
            ],
            "temperature": 0.7,
            "max_tokens": 4000
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(
                "https://api.openai.com/v1/chat/completions",
                headers=headers,
                json=data,
                timeout=120.0
            )
            response.raise_for_status()
            result = response.json()
            
            answer = result["choices"][0]["message"]["content"]
            
            # task2knowledgeと同じようにシンプルに処理
            return {
                "question": "プロンプト適用結果",
                "answer": answer.strip()
            }

    async def _call_anthropic(self, content: str, prompt: str, model: str) -> Dict[str, str]:
        headers = {
            "x-api-key": self.anthropic_api_key,
            "Content-Type": "application/json",
            "anthropic-version": "2023-06-01"
        }
        
        data = {
            "model": model or "claude-3-haiku-20240307",
            "max_tokens": 1000,
            "messages": [
                {"role": "user", "content": f"{prompt}\n\n{content}"}
            ]
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(
                "https://api.anthropic.com/v1/messages",
                headers=headers,
                json=data,
                timeout=60.0
            )
            response.raise_for_status()
            result = response.json()
            
            answer = result["content"][0]["text"]
            return self._parse_qa_response(answer)

    async def _call_gemini(self, content: str, prompt: str, model: str) -> Dict[str, str]:
        headers = {
            "Content-Type": "application/json"
        }
        
        data = {
            "contents": [{
                "parts": [{
                    "text": f"{prompt}\n\n{content}"
                }]
            }],
            "generationConfig": {
                "temperature": 0.7,
                "maxOutputTokens": 1000
            }
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(
                f"https://generativelanguage.googleapis.com/v1beta/models/{model or 'gemini-pro'}:generateContent?key={self.gemini_api_key}",
                headers=headers,
                json=data,
                timeout=60.0
            )
            response.raise_for_status()
            result = response.json()
            
            answer = result["candidates"][0]["content"]["parts"][0]["text"]
            return self._parse_qa_response(answer)

    async def _call_lmstudio(self, content: str, prompt: str, model: str) -> Dict[str, str]:
        headers = {"Content-Type": "application/json"}
        
        data = {
            "model": model,
            "messages": [
                {"role": "system", "content": prompt},
                {"role": "user", "content": content}
            ],
            "temperature": 0.7
        }
        
        async with httpx.AsyncClient() as client:
            response = await client.post(
                f"{self.lm_studio_base_url}/chat/completions",
                headers=headers,
                json=data,
                timeout=60.0
            )
            response.raise_for_status()
            result = response.json()
            
            answer = result["choices"][0]["message"]["content"]
            return self._parse_qa_response(answer)

    def _parse_qa_response(self, response: str) -> Dict[str, str]:
        """Parse AI response to extract question and answer"""
        lines = response.strip().split('\n')
        question = ""
        answer = ""
        
        current_section = None
        for line in lines:
            line = line.strip()
            if line.startswith("Q:") or line.startswith("Question:"):
                current_section = "question"
                question = line.replace("Q:", "").replace("Question:", "").strip()
            elif line.startswith("A:") or line.startswith("Answer:"):
                current_section = "answer"
                answer = line.replace("A:", "").replace("Answer:", "").strip()
            elif current_section == "question" and line:
                question += " " + line
            elif current_section == "answer" and line:
                answer += " " + line
        
        if not question and not answer:
            question = "Generated Knowledge"
            answer = response.strip()
        elif not question:
            question = "Generated Knowledge"
        elif not answer:
            answer = response.strip()
        
        return {"question": question.strip(), "answer": answer.strip()}

    def _is_test_mode(self) -> bool:
        """Check if we're in test mode (no valid API keys configured for default provider)"""
        # デフォルトプロバイダー（OpenAI）のAPIキーのみをチェック
        if self.default_provider == "openai":
            return not self.openai_api_key or self.openai_api_key == "your-openai-api-key"
        elif self.default_provider == "anthropic":
            return not self.anthropic_api_key or self.anthropic_api_key == "your-anthropic-api-key"
        elif self.default_provider == "gemini":
            return not self.gemini_api_key or self.gemini_api_key == "your-gemini-api-key"
        else:
            return True

    def _generate_test_response(self, content: str, prompt: str) -> Dict[str, str]:
        """Generate a test response when API keys are not available"""
        content_preview = content[:200] + "..." if len(content) > 200 else content
        
        if "summary" in prompt.lower() or "summarize" in prompt.lower():
            question = "What is the main summary of this content?"
            answer = f"This content discusses {content_preview}. The main points include key concepts and findings relevant to the subject matter."
        elif "evaluation" in prompt.lower() or "assess" in prompt.lower() or "rate" in prompt.lower():
            question = "What is the evaluation result?"
            answer = f"Based on the analysis criteria, this content scores well in several areas. Rating: 4/5. The content demonstrates good quality and meets most evaluation standards."
        elif "technical" in prompt.lower() or "methodology" in prompt.lower():
            question = "What are the technical aspects?"
            answer = f"The technical methodology involves systematic approaches and established practices. Key technical elements are well-documented and follow industry standards."
        elif "security" in prompt.lower() or "risk" in prompt.lower():
            question = "What are the security considerations?"
            answer = f"Security analysis indicates moderate risk level. Recommendations include standard security practices and regular monitoring."
        elif "innovation" in prompt.lower() or "future" in prompt.lower():
            question = "What are the innovation aspects?"
            answer = f"This represents innovative approaches with potential for future development. The concepts show promise for advancement in the field."
        else:
            question = f"Analysis based on: {prompt[:100]}..."
            answer = f"The analysis reveals important insights about {content_preview}. The findings suggest significant value and relevance to the specified criteria."
        
        return {"question": question, "answer": answer}

class MarkItDownService:
    def __init__(self):
        self.server_url = os.getenv("MARKITDOWN_SERVER_URL", "http://localhost:8001")
        self.api_key = os.getenv("MARKITDOWN_API_KEY", "your-api-key")

    async def convert_file(self, file_path: str) -> str:
        """Convert file to markdown using MarkItDown server"""
        try:
            async with httpx.AsyncClient() as client:
                with open(file_path, 'rb') as f:
                    files = {'file': f}
                    headers = {'X-API-Key': self.api_key}
                    response = await client.post(
                        f"{self.server_url}/convert",
                        files=files,
                        headers=headers,
                        timeout=120.0
                    )
                    response.raise_for_status()
                    result = response.json()
                    return result.get("markdown", "")
        except Exception as e:
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    return f.read()
            except:
                raise Exception(f"Failed to convert file: {str(e)}")

class TaskService:
    def __init__(self, db: Session):
        self.db = db
        self.ai_service = AIService()

    async def process_task(self, task_id: int):
        """Process a single task"""
        task = self.db.query(Task).filter(Task.id == task_id).first()
        if not task:
            return
        
        try:
            task.status = "processing"
            self.db.commit()
            
            record = self.db.query(Record).filter(Record.id == task.record_id).first()
            prompt = self.db.query(Prompt).filter(Prompt.id == task.prompt_id).first()
            
            if not record or not prompt:
                task.status = "failed"
                task.error_message = "Record or prompt not found"
                self.db.commit()
                return
            
            result = await self.ai_service.generate_knowledge(
                record.content,
                prompt.content,
                task.ai_provider,
                task.ai_model
            )
            
            knowledge = Knowledge(
                question=result["question"],
                answer=result["answer"],
                record_id=task.record_id,
                prompt_id=task.prompt_id,
                user_id=task.user_id,
                parent_type=record.parent_type,
                parent_id=record.parent_id
            )
            
            self.db.add(knowledge)
            self.db.flush()
            
            task.status = "completed"
            task.result_knowledge_id = knowledge.id
            self.db.commit()
            
            user = self.db.query(User).filter(User.id == task.user_id).first()
            if user:
                user.current_month_usage += 1
                self.db.commit()
                
        except Exception as e:
            task.status = "failed"
            task.error_message = str(e)
            self.db.commit()

class ExportService:
    def __init__(self, db: Session):
        self.db = db

    def generate_knowledge_matrix(self, user_id: int, group_ids: Optional[List[int]] = None) -> List[KnowledgeMatrixItem]:
        """Generate knowledge matrix data"""
        query = self.db.query(
            Knowledge, Record, Group, Prompt
        ).join(
            Record, Knowledge.record_id == Record.id
        ).join(
            Group, Record.group_id == Group.id
        ).join(
            Prompt, Knowledge.prompt_id == Prompt.id
        ).filter(
            Knowledge.user_id == user_id,
            Knowledge.deleted == False,
            Record.deleted == False,
            Group.deleted == False,
            Prompt.deleted == False
        )
        
        if group_ids:
            query = query.filter(Group.id.in_(group_ids))
        
        results = query.all()
        
        matrix_items = []
        for knowledge, record, group, prompt in results:
            matrix_items.append(KnowledgeMatrixItem(
                group_name=group.name,
                record_title=record.title,
                prompt_name=prompt.name,
                question=knowledge.question,
                answer=knowledge.answer,
                created_at=knowledge.created_at
            ))
        
        return matrix_items

    def export_to_excel(self, matrix_items: List[KnowledgeMatrixItem], file_path: str):
        """Export knowledge matrix to Excel"""
        data = []
        for item in matrix_items:
            data.append({
                "Group": item.group_name,
                "Record": item.record_title,
                "Prompt": item.prompt_name,
                "Question": item.question,
                "Answer": item.answer,
                "Created At": item.created_at.strftime("%Y-%m-%d %H:%M:%S")
            })
        
        df = pd.DataFrame(data)
        df.to_excel(file_path, index=False, engine='openpyxl')
        return file_path

class MatrixService:
    def __init__(self, db: Session):
        self.db = db
    
    def create_matrix(self, matrix_data: MatrixCreate, user_id: int) -> Matrix:
        """マトリックス定義を作成"""
        matrix = Matrix(
            name=matrix_data.name,
            description=matrix_data.description,
            group_ids=matrix_data.group_ids,
            user_id=user_id
        )
        self.db.add(matrix)
        self.db.commit()
        self.db.refresh(matrix)
        return matrix
    
    def update_matrix(self, matrix_id: int, matrix_data: MatrixUpdate, user_id: int) -> Matrix:
        """マトリックス定義を更新"""
        matrix = self.db.query(Matrix).filter(
            Matrix.id == matrix_id,
            Matrix.user_id == user_id,
            Matrix.deleted == False
        ).first()
        
        if not matrix:
            raise ValueError("Matrix not found")
        
        if matrix_data.name is not None:
            matrix.name = matrix_data.name
        if matrix_data.description is not None:
            matrix.description = matrix_data.description
        if matrix_data.group_ids is not None:
            matrix.group_ids = matrix_data.group_ids
        
        matrix.updated_at = datetime.utcnow()
        self.db.commit()
        self.db.refresh(matrix)
        return matrix
    
    def delete_matrix(self, matrix_id: int, user_id: int) -> None:
        """マトリックス定義を削除"""
        matrix = self.db.query(Matrix).filter(
            Matrix.id == matrix_id,
            Matrix.user_id == user_id,
            Matrix.deleted == False
        ).first()
        
        if not matrix:
            raise ValueError("Matrix not found")
        
        matrix.deleted = True
        matrix.updated_at = datetime.utcnow()
        self.db.commit()
    
    def generate_matrix_data(self, matrix_id: int, user_id: int, include_plain_knowledge: bool = False) -> MatrixExportData:
        """マトリックス定義に基づいてナレッジマトリックスデータを生成"""
        # マトリックス定義を取得
        matrix = self.db.query(Matrix).filter(
            Matrix.id == matrix_id,
            Matrix.user_id == user_id,
            Matrix.deleted == False
        ).first()
        
        if not matrix:
            raise ValueError("Matrix not found")
        
        # グループIDをパース
        group_ids = [int(id.strip()) for id in matrix.group_ids.split(',') if id.strip()]
        
        # グループ情報を取得（順序を保持）
        groups = []
        for group_id in group_ids:
            group = self.db.query(Group).filter(
                Group.id == group_id,
                Group.user_id == user_id,
                Group.deleted == False
            ).first()
            if group:
                groups.append(group)
        
        # 各グループに関連するプロンプトを取得
        group_prompts = {}
        for group in groups:
            prompts = self.db.query(Prompt).join(
                Knowledge, Knowledge.prompt_id == Prompt.id
            ).filter(
                Knowledge.record_id.in_(
                    self.db.query(Record.id).filter(
                        Record.group_id == group.id,
                        Record.deleted == False
                    )
                ),
                Knowledge.deleted == False,
                Prompt.deleted == False
            ).distinct().all()
            
            group_prompts[group.id] = {
                'name': group.name,
                'prompts': prompts
            }
        
        # プレーンナレッジを取得（必要な場合）
        plain_knowledge = {}
        if include_plain_knowledge and groups:
            first_group = groups[0]
            records = self.db.query(Record).filter(
                Record.group_id == first_group.id,
                Record.deleted == False
            ).all()
            for record in records:
                plain_knowledge[record.title] = record.content
        
        # レコードタイトルをユニークに取得
        record_titles = set()
        for group in groups:
            records = self.db.query(Record).filter(
                Record.group_id == group.id,
                Record.deleted == False
            ).all()
            for record in records:
                record_titles.add(record.title)
        
        record_titles = sorted(list(record_titles))
        
        # プロンプト名を取得
        all_prompts = set()
        for group_data in group_prompts.values():
            for prompt in group_data['prompts']:
                all_prompts.add(prompt.name)
        prompt_names = sorted(list(all_prompts))
        
        # ナレッジデータを生成
        knowledge_data = {}
        for record_title in record_titles:
            knowledge_data[record_title] = {}
            
            # このタイトルに一致するすべてのレコードのIDを取得
            record_ids = self.db.query(Record.id).filter(
                Record.title == record_title,
                Record.group_id.in_(group_ids),
                Record.deleted == False
            ).all()
            record_ids = [r[0] for r in record_ids]
            
            if record_ids:
                # このレコードに関連するナレッジを取得
                knowledge_records = self.db.query(Knowledge, Prompt).join(
                    Prompt, Knowledge.prompt_id == Prompt.id
                ).filter(
                    Knowledge.record_id.in_(record_ids),
                    Knowledge.deleted == False
                ).all()
                
                for knowledge, prompt in knowledge_records:
                    knowledge_data[record_title][prompt.name] = knowledge.answer
        
        return MatrixExportData(
            matrix_name=matrix.name,
            records=record_titles,
            prompts=prompt_names,
            knowledge_data=knowledge_data,
            plain_knowledge=plain_knowledge if include_plain_knowledge else None
        )
    
    def export_matrix_to_excel(self, matrix_data: MatrixExportData, file_path: str):
        """マトリックスデータをExcelファイルにエクスポート（LLMKnowledge2形式）"""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")
        
        wb = Workbook()
        ws = wb.active
        ws.title = "Knowledge Matrix"
        
        # ヘッダー行の作成
        ws['A1'] = 'ID'
        col = 2  # B列から開始
        for prompt_name in matrix_data.prompts:
            ws.cell(row=1, column=col, value=prompt_name)
            col += 1
        
        # プレーンナレッジ列を追加（必要な場合）
        if matrix_data.plain_knowledge:
            ws.cell(row=1, column=col, value='元テキスト')
        
        # データ行の配置
        row = 2
        for record_title in matrix_data.records:
            ws.cell(row=row, column=1, value=record_title)
            
            col = 2
            for prompt_name in matrix_data.prompts:
                value = matrix_data.knowledge_data.get(record_title, {}).get(prompt_name, '')
                ws.cell(row=row, column=col, value=value)
                col += 1
            
            # プレーンナレッジを追加（必要な場合）
            if matrix_data.plain_knowledge:
                plain_text = matrix_data.plain_knowledge.get(record_title, '')
                ws.cell(row=row, column=col, value=plain_text)
            
            row += 1
        
        # スタイルの設定
        last_col_num = len(matrix_data.prompts) + 1
        if matrix_data.plain_knowledge:
            last_col_num += 1
        
        # ヘッダー行のスタイル（中央揃え）
        for col in range(1, last_col_num + 1):
            cell = ws.cell(row=1, column=col)
            cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # データ行のスタイル（左上揃え）
        for row_num in range(2, row):
            for col in range(1, last_col_num + 1):
                cell = ws.cell(row=row_num, column=col)
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        
        # 列幅を設定（200px ≈ 28.571 units）
        for col in range(1, last_col_num + 1):
            ws.column_dimensions[chr(64 + col)].width = 28.571
        
        wb.save(file_path)

def log_history(db: Session, table_name: str, record_id: int, action: str, data: dict, user_id: Optional[int] = None):
    """Log changes to history table"""
    history = History(
        table_name=table_name,
        record_id=record_id,
        action=action,
        data=data,
        user_id=user_id
    )
    db.add(history)
    db.commit()
